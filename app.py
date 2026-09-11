import streamlit as st
import pandas as pd
import io
import json
import datetime
import google.generativeai as genai
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.worksheet.datavalidation import DataValidation

# =========================================================
# KONFIGURASI TAMPILAN WEB
# =========================================================
st.set_page_config(page_title="Dashboard Konveksi GK Garment Indonesia", layout="wide", page_icon="🧵")

st.title("🧵 Dashboard Produksi Konveksi GK Garment Indonesia")
st.write(
    "Upload file Excel pesanan customer, atur asumsi produksi, lalu biarkan AI merapikan, "
    "memvalidasi, dan menghitung rekap produksi — lengkap dengan rumus Excel yang bisa "
    "kamu ubah sendiri, dan dashboard ringkasan."
)

BASE_COLS = ["No", "Kode Pesanan", "Nama Item", "Warna", "Total Pcs", "Jenis Sablon/Bordir", "Prioritas", "Catatan", "Sumber Tab"]
REQUIRED_AI_KEYS = ["Kode Pesanan", "Nama Item", "Warna", "Total Pcs", "Jenis Sablon/Bordir", "Prioritas", "Catatan"]

# =========================================================
# SIDEBAR PENGATURAN
# =========================================================
with st.sidebar:
    st.header("⚙️ Pengaturan AI")
    api_key = st.text_input("Masukkan Gemini API Key:", type="password", help="Dapatkan gratis di aistudio.google.com")
    model_name = st.selectbox("Model Gemini", ["gemini-3.6-flash", "gemini-3.5-flash", "gemini-3.1-flash-lite"], index=0)

    st.markdown("---")
    st.header("🧮 Asumsi Perhitungan")
    rasio_kain = st.number_input("Rasio Kain (pcs per kg)", min_value=0.1, value=3.5, step=0.1,
                                  help="Berapa pcs pakaian yang bisa dibuat dari 1 kg kain.")
    harga_kain = st.number_input("Harga Kain per Kg (Rp)", min_value=0, value=45000, step=1000)
    prioritas_tinggi = st.number_input("Batas pcs untuk Prioritas Tinggi", min_value=1, value=300, step=10)
    prioritas_sedang = st.number_input("Batas pcs untuk Prioritas Sedang", min_value=1, value=100, step=10)

    st.markdown("---")
    st.markdown("**Cara Pakai:**")
    st.markdown("1. Masukkan API Key & atur asumsi di atas")
    st.markdown("2. Upload file Excel pesanan")
    st.markdown("3. Sesuaikan instruksi/prompt bila perlu")
    st.markdown("4. Klik **Olah Data Sekarang**")

if api_key:
    genai.configure(api_key=api_key)

# =========================================================
# PROMPT DEFAULT — DIPERKUAT
# =========================================================
default_prompt = f"""Kamu adalah admin produksi konveksi senior yang teliti. Olah data pesanan berikut menjadi
satu tabel rekap produksi yang BERSIH dan SIAP PAKAI, dengan langkah wajib berikut:

1. STANDARISASI: samakan penulisan Nama Item dan Warna yang sebenarnya sama tapi ditulis beda
   (contoh: "kaos polo", "Kaos Polo ", "KAOS POLO" -> "Kaos Polo"). Gunakan Title Case.
2. GABUNGKAN DUPLIKAT: jika ada baris dengan Kode Pesanan + Nama Item + Warna yang identik
   setelah distandarisasi, jumlahkan Total Pcs-nya menjadi satu baris saja.
3. VALIDASI: jika Total Pcs kosong, 0, negatif, atau tidak berupa angka, tetap sertakan barisnya
   dengan Total Pcs = 0 dan jelaskan masalahnya di kolom Catatan (mis. "Qty tidak terbaca dari sumber").
4. TENTUKAN PRIORITAS setiap baris berdasarkan Total Pcs:
   - Total Pcs >= {int(prioritas_tinggi)} -> "Tinggi"
   - {int(prioritas_sedang)} <= Total Pcs < {int(prioritas_tinggi)} -> "Sedang"
   - Total Pcs < {int(prioritas_sedang)} -> "Rendah"
   Jika data asli punya kolom tanggal/deadline, pertimbangkan juga: deadline < 7 hari dari sekarang
   otomatis "Tinggi" berapa pun jumlah pcs-nya, dan sebutkan alasannya singkat di Catatan.
5. CATATAN: isi HANYA jika ada hal yang perlu diperhatikan admin (data ambigu, digabung dari baris
   lain, quantity mencurigakan, deadline mepet). Kosongkan ("") jika tidak ada catatan.
6. JANGAN menghitung estimasi kain atau biaya — itu akan dihitung otomatis oleh rumus Excel.

"Kembalikan HANYA array JSON murni (tanpa pembungkus markdown ```json, tanpa teks pengantar, tanpa penjelasan penutup) yang valid.".
"""

# =========================================================
# FUNGSI: BANGUN WORKBOOK EXCEL (rumus + styling)
# =========================================================
def build_excel_workbook(df: pd.DataFrame, rasio_kain: float, harga_kain: float) -> io.BytesIO:
    HEADER_FILL = PatternFill("solid", fgColor="1F3864")
    HEADER_FONT = Font(color="FFFFFF", bold=True, name="Arial", size=11)
    ASSUMPTION_FILL = PatternFill("solid", fgColor="FFF2CC")
    ASSUMPTION_FONT = Font(color="0000FF", bold=True, name="Arial", size=11)
    TITLE_FONT = Font(bold=True, size=16, name="Arial", color="1F3864")
    TOTAL_FILL = PatternFill("solid", fgColor="D9E1F2")
    TOTAL_FONT = Font(bold=True, name="Arial")
    BAND_FILL = PatternFill("solid", fgColor="F2F2F2")
    THIN = Side(style="thin", color="BFBFBF")
    BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    PRIORITY_COLORS = {
        "Tinggi": PatternFill("solid", fgColor="F8CBAD"),
        "Sedang": PatternFill("solid", fgColor="FFE699"),
        "Rendah": PatternFill("solid", fgColor="C6E0B4"),
    }

    wb = Workbook()
    ws = wb.active
    ws.title = "Rekap Produksi"

    ws.merge_cells("A1:J1")
    ws["A1"] = "REKAP PRODUKSI KONVEKSI"
    ws["A1"].font = TITLE_FONT
    ws.row_dimensions[1].height = 26
    ws["A2"] = f"Dibuat otomatis: {datetime.datetime.now().strftime('%d %B %Y %H:%M')}"
    ws["A2"].font = Font(italic=True, size=9, color="808080")

    # Blok asumsi (input yang boleh diubah user -> font biru + fill kuning)
    ws["A4"] = "Asumsi Rasio Kain (pcs/kg)"
    ws["B4"] = rasio_kain
    ws["A5"] = "Asumsi Harga Kain (Rp/kg)"
    ws["B5"] = harga_kain
    for cell in ["A4", "A5"]:
        ws[cell].font = Font(bold=True, name="Arial")
    for cell in ["B4", "B5"]:
        ws[cell].fill = ASSUMPTION_FILL
        ws[cell].font = ASSUMPTION_FONT
        ws[cell].border = BORDER
    ws["B4"].number_format = "0.00"
    ws["B5"].number_format = '"Rp"#,##0'
    RATIO_CELL, HARGA_CELL = "$B$4", "$B$5"

    base_cols = BASE_COLS
    extra_cols = ["Estimasi Kain (Kg)", "Estimasi Biaya Kain (Rp)"]
    all_cols = base_cols + extra_cols
    header_row = 7
    for idx, col_name in enumerate(all_cols, start=1):
        c = ws.cell(row=header_row, column=idx, value=col_name)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.border = BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[header_row].height = 32

    pcs_col_idx = base_cols.index("Total Pcs") + 1
    pcs_col_letter = get_column_letter(pcs_col_idx)
    prioritas_col_idx = base_cols.index("Prioritas") + 1
    prioritas_col_letter = get_column_letter(prioritas_col_idx)
    jenis_col_idx = base_cols.index("Jenis Sablon/Bordir") + 1
    jenis_col_letter = get_column_letter(jenis_col_idx)
    kain_col_idx = len(base_cols) + 1
    biaya_col_idx = len(base_cols) + 2
    kain_col_letter = get_column_letter(kain_col_idx)
    biaya_col_letter = get_column_letter(biaya_col_idx)

    first_data_row = header_row + 1
    n_rows = len(df)
    for i, (_, row) in enumerate(df.iterrows()):
        r = first_data_row + i
        for j, col_name in enumerate(base_cols, start=1):
            if col_name == "No":
                val = i + 1
            else:
                val = row.get(col_name, "")
            cell = ws.cell(row=r, column=j, value=val)
            cell.border = BORDER
            if i % 2 == 1:
                cell.fill = BAND_FILL
        kain_cell = ws.cell(row=r, column=kain_col_idx, value=f"={pcs_col_letter}{r}/{RATIO_CELL}")
        kain_cell.number_format = "0.00"
        kain_cell.border = BORDER
        biaya_cell = ws.cell(row=r, column=biaya_col_idx, value=f"={kain_col_letter}{r}*{HARGA_CELL}")
        biaya_cell.number_format = '"Rp"#,##0'
        biaya_cell.border = BORDER
        if i % 2 == 1:
            kain_cell.fill = BAND_FILL
            biaya_cell.fill = BAND_FILL

    last_data_row = first_data_row + n_rows - 1 if n_rows > 0 else first_data_row

    total_row = last_data_row + 1
    ws.cell(row=total_row, column=1, value="TOTAL").font = TOTAL_FONT
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=pcs_col_idx - 1)
    ws.cell(row=total_row, column=1).alignment = Alignment(horizontal="right")

    if n_rows > 0:
        total_pcs_formula = f"=SUM({pcs_col_letter}{first_data_row}:{pcs_col_letter}{last_data_row})"
        total_kain_formula = f"=SUM({kain_col_letter}{first_data_row}:{kain_col_letter}{last_data_row})"
        total_biaya_formula = f"=SUM({biaya_col_letter}{first_data_row}:{biaya_col_letter}{last_data_row})"
    else:
        total_pcs_formula = total_kain_formula = total_biaya_formula = 0

    ws.cell(row=total_row, column=pcs_col_idx, value=total_pcs_formula).number_format = "#,##0"
    ws.cell(row=total_row, column=kain_col_idx, value=total_kain_formula).number_format = "0.00"
    ws.cell(row=total_row, column=biaya_col_idx, value=total_biaya_formula).number_format = '"Rp"#,##0'

    for col in range(1, len(all_cols) + 1):
        c = ws.cell(row=total_row, column=col)
        c.fill = TOTAL_FILL
        c.font = TOTAL_FONT
        c.border = Border(top=Side(style="double"), bottom=THIN, left=THIN, right=THIN)

    if n_rows > 0:
        prio_range = f"{prioritas_col_letter}{first_data_row}:{prioritas_col_letter}{last_data_row}"
        for label in ["Tinggi", "Sedang", "Rendah"]:
            ws.conditional_formatting.add(
                prio_range, CellIsRule(operator="equal", formula=[f'"{label}"'], fill=PRIORITY_COLORS[label])
            )
        dv = DataValidation(type="list", formula1='"Tinggi,Sedang,Rendah"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(prio_range)

        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(all_cols))}{last_data_row}"
    ws.freeze_panes = f"A{first_data_row}"

    widths = [5, 14, 20, 14, 11, 18, 11, 30, 14, 16, 20]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # ================= SHEET RINGKASAN =================
    ws2 = wb.create_sheet("Ringkasan")
    ws2.merge_cells("A1:D1")
    ws2["A1"] = "RINGKASAN PRODUKSI"
    ws2["A1"].font = TITLE_FONT
    ws2.row_dimensions[1].height = 26

    kode_col_letter = get_column_letter(base_cols.index("Kode Pesanan") + 1)
    if n_rows > 0:
        total_pesanan_formula = (
            f"=SUMPRODUCT(1/COUNTIF('Rekap Produksi'!{kode_col_letter}{first_data_row}:"
            f"{kode_col_letter}{last_data_row},'Rekap Produksi'!{kode_col_letter}{first_data_row}:"
            f"{kode_col_letter}{last_data_row}))"
        )
    else:
        total_pesanan_formula = 0

    kpi_labels = [
        ("Total Pesanan (Kode Unik)", total_pesanan_formula, "#,##0"),
        ("Total Pcs Keseluruhan", f"='Rekap Produksi'!{pcs_col_letter}{total_row}", "#,##0"),
        ("Total Estimasi Kain (Kg)", f"='Rekap Produksi'!{kain_col_letter}{total_row}", "0.00"),
        ("Total Estimasi Biaya Kain (Rp)", f"='Rekap Produksi'!{biaya_col_letter}{total_row}", '"Rp"#,##0'),
    ]
    r = 3
    for label, formula, fmt in kpi_labels:
        ws2.cell(row=r, column=1, value=label).font = Font(bold=True, name="Arial")
        ws2.cell(row=r, column=1).border = BORDER
        val_cell = ws2.cell(row=r, column=2, value=formula)
        val_cell.number_format = fmt
        val_cell.fill = PatternFill("solid", fgColor="D9E1F2")
        val_cell.font = Font(bold=True, size=12, name="Arial")
        val_cell.border = BORDER
        r += 1

    r += 1
    ws2.cell(row=r, column=1, value="Breakdown per Jenis Sablon/Bordir").font = Font(bold=True, size=12, name="Arial")
    r += 1
    breakdown_header_row = r
    for j, col_name in enumerate(["Jenis", "Total Pcs"], start=1):
        c = ws2.cell(row=r, column=j, value=col_name)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.border = BORDER
    r += 1
    breakdown_first_row = r
    jenis_unik = sorted(df["Jenis Sablon/Bordir"].dropna().astype(str).unique().tolist()) if n_rows > 0 else []
    for jenis in jenis_unik:
        ws2.cell(row=r, column=1, value=jenis).border = BORDER
        f = (f"=SUMIF('Rekap Produksi'!{jenis_col_letter}{first_data_row}:{jenis_col_letter}{last_data_row},"
             f"A{r},'Rekap Produksi'!{pcs_col_letter}{first_data_row}:{pcs_col_letter}{last_data_row})")
        fc = ws2.cell(row=r, column=2, value=f)
        fc.number_format = "#,##0"
        fc.border = BORDER
        r += 1
    breakdown_last_row = r - 1

    # Breakdown per Prioritas (COUNTIF)
    r += 1
    ws2.cell(row=r, column=1, value="Jumlah Pesanan per Prioritas").font = Font(bold=True, size=12, name="Arial")
    r += 1
    prio_header_row = r
    for j, col_name in enumerate(["Prioritas", "Jumlah Baris"], start=1):
        c = ws2.cell(row=r, column=j, value=col_name)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.border = BORDER
    r += 1
    prio_first_row = r
    for label in ["Tinggi", "Sedang", "Rendah"]:
        ws2.cell(row=r, column=1, value=label).border = BORDER
        f = (f"=COUNTIF('Rekap Produksi'!{prioritas_col_letter}{first_data_row}:"
             f"{prioritas_col_letter}{last_data_row},A{r})") if n_rows > 0 else 0
        fc = ws2.cell(row=r, column=2, value=f)
        fc.number_format = "#,##0"
        fc.border = BORDER
        r += 1
    prio_last_row = r - 1

    ws2.column_dimensions["A"].width = 32
    ws2.column_dimensions["B"].width = 22

    if jenis_unik:
        chart = BarChart()
        chart.title = "Total Pcs per Jenis Sablon/Bordir"
        chart.y_axis.title = "Pcs"
        chart.x_axis.title = "Jenis"
        data_ref = Reference(ws2, min_col=2, min_row=breakdown_header_row, max_row=breakdown_last_row)
        cats_ref = Reference(ws2, min_col=1, min_row=breakdown_first_row, max_row=breakdown_last_row)
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats_ref)
        chart.width, chart.height = 16, 8
        ws2.add_chart(chart, "D3")

        pie = PieChart()
        pie.title = "Proporsi Prioritas Pesanan"
        pdata = Reference(ws2, min_col=2, min_row=prio_header_row, max_row=prio_last_row)
        pcats = Reference(ws2, min_col=1, min_row=prio_first_row, max_row=prio_last_row)
        pie.add_data(pdata, titles_from_data=True)
        pie.set_categories(pcats)
        pie.width, pie.height = 16, 8
        ws2.add_chart(pie, "D19")

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


# =========================================================
# AREA UPLOAD FILE
# =========================================================
uploaded_file = st.file_uploader("📂 Upload File Excel Pesanan Customer (.xlsx / .xls)", type=["xlsx", "xls"])

if uploaded_file:
    try:
        xls = pd.ExcelFile(uploaded_file)
        all_sheets = []
        for sheet_name in xls.sheet_names:
            df_sheet = pd.read_excel(xls, sheet_name=sheet_name)
            df_sheet["Nama Tab Excel"] = sheet_name
            all_sheets.append(df_sheet)
        df_raw = pd.concat(all_sheets, ignore_index=True)

        st.write("### 📄 Preview Data Asli dari Customer")
        st.dataframe(df_raw, use_container_width=True)

        st.write("---")
        st.write("### 🤖 Instruksi Pengolahan Data (Prompt Admin)")
        user_prompt = st.text_area("Tuliskan instruksi/prompt untuk AI di sini:", value=default_prompt, height=320)

        if st.button("🚀 Olah Data Sekarang", type="primary"):
            if not api_key:
                st.error("⚠️ Silakan masukkan Gemini API Key terlebih dahulu di menu sidebar sebelah kiri!")
            else:
                with st.spinner("AI sedang membaca, membersihkan, dan mengolah data pesanan..."):
                    json_input = df_raw.to_json(orient="records", force_ascii=False)

                    system_prompt = f"""
                    Berikut data pesanan customer dari Excel (JSON):
                    {json_input}

                    Instruksi dari Admin:
                    {user_prompt}
                    """

                    try:
                        model = genai.GenerativeModel(model_name)
                        response = model.generate_content(system_prompt)
                    except Exception as api_err:
                        st.error(f"Gagal menghubungi Gemini API: {api_err}")
                        st.stop()

                    raw_text = response.text.strip()
                    for fence in ("```json", "```"):
                        if raw_text.startswith(fence):
                            raw_text = raw_text[len(fence):]
                    if raw_text.endswith("```"):
                        raw_text = raw_text[:-3]
                    raw_text = raw_text.strip()

                    try:
                        parsed_json = json.loads(raw_text)
                        if isinstance(parsed_json, dict):
                            parsed_json = parsed_json.get("data", [parsed_json])
                        df_result = pd.DataFrame(parsed_json)

                        # === TAMBAHKAN LOGIKA PEMETAAN FLEKSIBEL INI ===
                        # Cek apakah AI menggunakan nama kolom kuantitas yang berbeda (Qty / quantity / total_pcs)
                        qty_cols = [c for c in df_result.columns if c.lower() in ['qty', 'quantity', 'jumlah', 'total_pcs', 'total pcs']]
                        if "Total Pcs" not in df_result.columns and qty_cols:
                             df_result["Total Pcs"] = df_result[qty_cols[0]]

                        # Fallback Aman: Jika AI gagal mengembalikan angka Qty, ambil langsung dari data asli (df_raw)
                          if "Total Pcs" not in df_result.columns or df_result["Total Pcs"].sum() == 0:
                          if "Qty" in df_raw.columns:
                            df_result["Total Pcs"] = df_raw["Qty"]

                        # Validasi & pembersihan minimal terhadap hasil AI
                        for key in REQUIRED_AI_KEYS:
                            if key not in df_result.columns:
                                df_result[key] = "" if key != "Total Pcs" else 0
                        df_result["Total Pcs"] = pd.to_numeric(df_result["Total Pcs"], errors="coerce").fillna(0).astype(int)
                        if "Sumber Tab" not in df_result.columns:
                            df_result["Sumber Tab"] = ""
                        df_result["No"] = range(1, len(df_result) + 1)

                        st.success(f"✅ Data Berhasil Diolah oleh AI! ({len(df_result)} baris rekap)")

                        col1, col2, col3, col4 = st.columns(4)
                        col1.metric("Total Pesanan Unik", df_result["Kode Pesanan"].nunique())
                        col2.metric("Total Pcs", int(df_result["Total Pcs"].sum()))
                        col3.metric("Estimasi Kain (Kg)", f"{df_result['Total Pcs'].sum() / rasio_kain:,.2f}")
                        col4.metric("Estimasi Biaya Kain", f"Rp {df_result['Total Pcs'].sum() / rasio_kain * harga_kain:,.0f}")

                        st.write("### 📊 Hasil Rekapitulasi Produksi")
                        st.dataframe(df_result, use_container_width=True)

                        excel_bytes = build_excel_workbook(df_result, rasio_kain, harga_kain)

                        st.download_button(
                            label="📥 Download Hasil Olah Data (.xlsx)",
                            data=excel_bytes,
                            file_name=f"Rekap_Produksi_Konveksi_{datetime.date.today().isoformat()}.xlsx",
                            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        )
                    except Exception as json_err:
                        st.error(f"Gagal membaca respon AI sebagai tabel ({json_err}). Respon mentah dari AI:")
                        st.code(response.text)

    except Exception as e:
        st.error(f"Gagal membaca file Excel: {e}")
else:
    st.info("👆 Silakan upload file Excel pesanan customer untuk memulai.")
